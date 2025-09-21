#!/usr/bin/env python3
"""
Face Recognition Web Application PowerPoint Presentation Generator
Creates a comprehensive presentation based on the application structure and features.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement, qn

def create_face_recognition_presentation():
    """Create a comprehensive PowerPoint presentation for the Face Recognition Web Application."""
    
    # Create presentation object
    prs = Presentation()
    
    # Set slide dimensions (16:9 aspect ratio)
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    # Define color scheme
    primary_color = RGBColor(52, 152, 219)    # Blue
    secondary_color = RGBColor(26, 188, 156)  # Teal
    accent_color = RGBColor(231, 76, 60)      # Red
    text_color = RGBColor(44, 62, 80)         # Dark Blue
    
    # ===== SLIDE 1: Title Slide =====
    slide1_layout = prs.slide_layouts[0]  # Title slide layout
    slide1 = prs.slides.add_slide(slide1_layout)
    
    title1 = slide1.shapes.title
    subtitle1 = slide1.placeholders[1]
    
    title1.text = "Face Recognition Web Application"
    title1.text_frame.paragraphs[0].font.size = Pt(44)
    title1.text_frame.paragraphs[0].font.color.rgb = primary_color
    title1.text_frame.paragraphs[0].font.bold = True
    
    subtitle1.text = "Automated Attendance Management System\nUsing Computer Vision and Machine Learning\n\nDeveloped with Flask Framework"
    for paragraph in subtitle1.text_frame.paragraphs:
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = text_color
        paragraph.alignment = PP_ALIGN.CENTER
    
    # ===== SLIDE 2: System Overview & Architecture =====
    slide2_layout = prs.slide_layouts[1]  # Title and content layout
    slide2 = prs.slides.add_slide(slide2_layout)
    
    title2 = slide2.shapes.title
    title2.text = "System Overview & Architecture"
    title2.text_frame.paragraphs[0].font.size = Pt(32)
    title2.text_frame.paragraphs[0].font.color.rgb = primary_color
    title2.text_frame.paragraphs[0].font.bold = True
    
    # Add content textbox
    left = Inches(1)
    top = Inches(1.5)
    width = Inches(11)
    height = Inches(5)
    
    textbox2 = slide2.shapes.add_textbox(left, top, width, height)
    text_frame2 = textbox2.text_frame
    text_frame2.word_wrap = True
    
    # System Architecture points
    content2 = [
        "🏗️ Three-Tier Architecture:",
        "   • Presentation Layer: HTML/CSS/JavaScript Frontend",
        "   • Business Logic: Flask Web Framework (Python)",
        "   • Data Layer: SQLAlchemy ORM with Database",
        "",
        "👥 User Role Management:",
        "   • Developer: System administration and user registration",
        "   • Teacher: Attendance session management and reporting",
        "   • Student: Face recognition-based attendance marking",
        "",
        "🔧 Core Components:",
        "   • Face Detection: OpenCV Haar Cascade Classifier",
        "   • Face Recognition: LBPH (Local Binary Pattern Histogram) Algorithm",
        "   • Session Management: Flask-Login authentication",
        "   • Database: SQLite with SQLAlchemy ORM"
    ]
    
    for item in content2:
        p = text_frame2.add_paragraph()
        p.text = item
        if item.startswith("🏗️") or item.startswith("👥") or item.startswith("🔧"):
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = secondary_color
        elif item.startswith("   •"):
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
            p.level = 1
        else:
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
    
    # ===== SLIDE 3: Features & Functionality =====
    slide3_layout = prs.slide_layouts[1]  # Title and content layout
    slide3 = prs.slides.add_slide(slide3_layout)
    
    title3 = slide3.shapes.title
    title3.text = "Features & Functionality"
    title3.text_frame.paragraphs[0].font.size = Pt(32)
    title3.text_frame.paragraphs[0].font.color.rgb = primary_color
    title3.text_frame.paragraphs[0].font.bold = True
    
    # Add content textbox
    textbox3 = slide3.shapes.add_textbox(left, top, width, height)
    text_frame3 = textbox3.text_frame
    text_frame3.word_wrap = True
    
    # Features content
    content3 = [
        "📸 Face Recognition System:",
        "   • Real-time camera integration for photo capture",
        "   • Automated face detection using computer vision",
        "   • Training data generation from multiple photo samples",
        "   • High-accuracy face recognition for attendance",
        "",
        "👨‍🏫 Teacher Dashboard:",
        "   • Start/Stop attendance sessions by batch and course",
        "   • Real-time session monitoring and management",
        "   • Comprehensive attendance reports and analytics",
        "   • Course and student management interface",
        "",
        "🎓 Student Interface:",
        "   • Simple one-click attendance marking",
        "   • Real-time feedback during recognition process",
        "   • Personal attendance history and statistics",
        "   • Secure profile management with photo upload",
        "",
        "⚙️ Developer Tools:",
        "   • User registration and management system",
        "   • Database administration and maintenance",
        "   • System configuration and monitoring tools"
    ]
    
    for item in content3:
        p = text_frame3.add_paragraph()
        p.text = item
        if item.startswith("📸") or item.startswith("👨‍🏫") or item.startswith("🎓") or item.startswith("⚙️"):
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = secondary_color
        elif item.startswith("   •"):
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
            p.level = 1
        else:
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
    
    # ===== SLIDE 4: Technology Stack & Implementation =====
    slide4_layout = prs.slide_layouts[1]  # Title and content layout
    slide4 = prs.slides.add_slide(slide4_layout)
    
    title4 = slide4.shapes.title
    title4.text = "Technology Stack & Implementation"
    title4.text_frame.paragraphs[0].font.size = Pt(32)
    title4.text_frame.paragraphs[0].font.color.rgb = primary_color
    title4.text_frame.paragraphs[0].font.bold = True
    
    # Add content textbox
    textbox4 = slide4.shapes.add_textbox(left, top, width, height)
    text_frame4 = textbox4.text_frame
    text_frame4.word_wrap = True
    
    # Technology stack content
    content4 = [
        "🐍 Backend Technologies:",
        "   • Python 3.x - Core programming language",
        "   • Flask - Lightweight web framework",
        "   • SQLAlchemy ORM - Database abstraction layer",
        "   • Flask-Login - User session management",
        "   • Flask-Migrate - Database migration tool",
        "",
        "🤖 Computer Vision & AI:",
        "   • OpenCV - Computer vision library for image processing",
        "   • Haar Cascade Classifier - Face detection algorithm",
        "   • LBPH Face Recognizer - Local Binary Pattern Histogram",
        "   • NumPy - Numerical computing for image arrays",
        "",
        "🌐 Frontend Technologies:",
        "   • HTML5 - Semantic markup structure",
        "   • CSS3 - Responsive design and styling",
        "   • JavaScript - Dynamic user interactions",
        "   • WebRTC - Real-time camera access and streaming",
        "",
        "🔒 Security & Deployment:",
        "   • HTTPS SSL/TLS encryption with custom certificates",
        "   • Role-based access control and authentication",
        "   • Secure file upload and photo storage system",
        "   • Cross-platform deployment with network accessibility"
    ]
    
    for item in content4:
        p = text_frame4.add_paragraph()
        p.text = item
        if item.startswith("🐍") or item.startswith("🤖") or item.startswith("🌐") or item.startswith("🔒"):
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = secondary_color
        elif item.startswith("   •"):
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
            p.level = 1
        else:
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
    
    # ===== SLIDE 5: System Workflow & Process =====
    slide5_layout = prs.slide_layouts[1]  # Title and content layout
    slide5 = prs.slides.add_slide(slide5_layout)
    
    title5 = slide5.shapes.title
    title5.text = "System Workflow & Process"
    title5.text_frame.paragraphs[0].font.size = Pt(32)
    title5.text_frame.paragraphs[0].font.color.rgb = primary_color
    title5.text_frame.paragraphs[0].font.bold = True
    
    # Add content textbox
    textbox5 = slide5.shapes.add_textbox(left, top, width, height)
    text_frame5 = textbox5.text_frame
    text_frame5.word_wrap = True
    
    # Workflow content
    content5 = [
        "📋 Student Registration Process:",
        "   1. Developer registers student with personal details",
        "   2. Student captures multiple photo samples for training",
        "   3. System trains face recognition model using LBPH algorithm",
        "   4. Student profile is activated for attendance marking",
        "",
        "📚 Attendance Session Management:",
        "   1. Teacher logs in and starts attendance session",
        "   2. System creates session with batch, course, and time details",
        "   3. Students use face recognition to mark attendance",
        "   4. Teacher monitors real-time attendance status",
        "   5. Session ends with comprehensive attendance report",
        "",
        "🎯 Face Recognition Process:",
        "   1. Student accesses camera through web interface",
        "   2. System captures real-time video frames",
        "   3. OpenCV detects faces in the captured frames",
        "   4. LBPH algorithm compares detected face with trained model",
        "   5. Attendance is marked if confidence score meets threshold",
        "   6. Immediate feedback provided to student and teacher"
    ]
    
    for item in content5:
        p = text_frame5.add_paragraph()
        p.text = item
        if item.startswith("📋") or item.startswith("📚") or item.startswith("🎯"):
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = secondary_color
        elif item.strip().startswith(("1.", "2.", "3.", "4.", "5.", "6.")):
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
            p.level = 1
        else:
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
    
    # ===== SLIDE 6: Conclusion & Future Enhancements =====
    slide6_layout = prs.slide_layouts[1]  # Title and content layout
    slide6 = prs.slides.add_slide(slide6_layout)
    
    title6 = slide6.shapes.title
    title6.text = "Conclusion & Future Enhancements"
    title6.text_frame.paragraphs[0].font.size = Pt(32)
    title6.text_frame.paragraphs[0].font.color.rgb = primary_color
    title6.text_frame.paragraphs[0].font.bold = True
    
    # Add content textbox
    textbox6 = slide6.shapes.add_textbox(left, top, width, height)
    text_frame6 = textbox6.text_frame
    text_frame6.word_wrap = True
    
    # Conclusion content
    content6 = [
        "✅ Key Achievements:",
        "   • Automated attendance system reducing manual effort",
        "   • High accuracy face recognition with real-time processing",
        "   • Secure multi-user role-based access control",
        "   • Scalable web-based architecture for easy deployment",
        "",
        "🚀 Future Enhancements:",
        "   • Deep learning models (CNN) for improved accuracy",
        "   • Mobile application for enhanced accessibility",
        "   • Advanced analytics and reporting dashboard",
        "   • Integration with existing educational management systems",
        "   • Multi-factor authentication for enhanced security",
        "",
        "🎯 Impact & Benefits:",
        "   • Eliminates proxy attendance and manual errors",
        "   • Saves time for both teachers and students",
        "   • Provides accurate attendance analytics and insights",
        "   • Modernizes traditional attendance systems",
        "",
        "📞 Thank You for Your Attention!"
    ]
    
    for item in content6:
        p = text_frame6.add_paragraph()
        p.text = item
        if item.startswith("✅") or item.startswith("🚀") or item.startswith("🎯"):
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = secondary_color
        elif item.startswith("📞"):
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = accent_color
            p.alignment = PP_ALIGN.CENTER
        elif item.startswith("   •"):
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
            p.level = 1
        else:
            p.font.size = Pt(14)
            p.font.color.rgb = text_color
    
    # Save the presentation
    presentation_path = "/home/runner/work/FaceMate/FaceMate/Face_Recognition_Web_Application_Presentation.pptx"
    prs.save(presentation_path)
    
    print(f"✅ PowerPoint presentation created successfully!")
    print(f"📁 File saved at: {presentation_path}")
    print(f"📊 Total slides: {len(prs.slides)}")
    
    return presentation_path

if __name__ == "__main__":
    create_face_recognition_presentation()
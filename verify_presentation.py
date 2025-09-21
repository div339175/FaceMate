#!/usr/bin/env python3
"""
Verification script to check the contents of the generated PowerPoint presentation.
"""

from pptx import Presentation
import os

def verify_presentation():
    """Verify the contents and structure of the generated presentation."""
    
    presentation_path = "/home/runner/work/FaceMate/FaceMate/Face_Recognition_Web_Application_Presentation.pptx"
    
    if not os.path.exists(presentation_path):
        print("❌ Presentation file not found!")
        return False
    
    try:
        # Load the presentation
        prs = Presentation(presentation_path)
        
        print("🔍 PowerPoint Presentation Verification")
        print("=" * 50)
        print(f"📁 File: {presentation_path}")
        print(f"📊 Total Slides: {len(prs.slides)}")
        print(f"📐 Dimensions: {prs.slide_width.inches:.2f}\" x {prs.slide_height.inches:.2f}\"")
        print()
        
        # Check each slide
        for i, slide in enumerate(prs.slides, 1):
            print(f"📄 Slide {i}:")
            
            # Get slide title
            title_text = "No title"
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    if hasattr(shape, "text_frame") and shape.text_frame:
                        title_text = shape.text.strip()
                        break
            
            print(f"   📋 Title: {title_text}")
            
            # Count text shapes
            text_shapes = 0
            total_text_length = 0
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_shapes += 1
                    total_text_length += len(shape.text)
            
            print(f"   📝 Text Elements: {text_shapes}")
            print(f"   📏 Content Length: {total_text_length} characters")
            print()
        
        print("✅ Presentation verification completed successfully!")
        print(f"💾 File size: {os.path.getsize(presentation_path)} bytes")
        
        return True
        
    except Exception as e:
        print(f"❌ Error verifying presentation: {str(e)}")
        return False

if __name__ == "__main__":
    verify_presentation()
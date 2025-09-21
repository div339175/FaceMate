#!/usr/bin/env python3
"""
Extract and display the content of the PowerPoint presentation for preview.
"""

from pptx import Presentation

def display_presentation_content():
    """Display the full content of the presentation slides."""
    
    presentation_path = "/home/runner/work/FaceMate/FaceMate/Face_Recognition_Web_Application_Presentation.pptx"
    
    try:
        prs = Presentation(presentation_path)
        
        print("🎯 FACE RECOGNITION WEB APPLICATION PRESENTATION")
        print("=" * 80)
        print()
        
        for i, slide in enumerate(prs.slides, 1):
            print(f"📄 SLIDE {i}")
            print("-" * 40)
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    # Clean up the text and format it nicely
                    text = shape.text.strip()
                    if text:
                        lines = text.split('\n')
                        for line in lines:
                            if line.strip():
                                print(line)
                        print()
            
            print("=" * 80)
            print()
        
        print("✅ Presentation content displayed successfully!")
        
    except Exception as e:
        print(f"❌ Error reading presentation: {str(e)}")

if __name__ == "__main__":
    display_presentation_content()